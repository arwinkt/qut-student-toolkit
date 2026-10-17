"""Real local OCR acceptance using generated text, never accounts or teaching files."""
from pathlib import Path
from tempfile import TemporaryDirectory
from qut_toolkit.local_ocr import extract_local_ocr, INCOMPLETE


def main():
    from PIL import Image, ImageDraw, ImageFont
    import pymupdf
    from docx import Document
    from pptx import Presentation
    from pptx.util import Inches
    with TemporaryDirectory(prefix='qut-ocr-smoke-') as directory:
        root=Path(directory).resolve()
        image=root/'synthetic.png'
        with Image.new('RGB',(800,240),'white') as canvas:
            draw=ImageDraw.Draw(canvas);font=ImageFont.load_default(size=42)
            draw.text((30,35),'COURSE WEEK 12',font=font,fill='black')
            draw.text((30,110),'REVISION QUESTIONS',font=font,fill='black')
            canvas.save(image)
        pdf=root/'synthetic.pdf'
        with pymupdf.open() as document:
            page=document.new_page(width=800,height=240)
            page.insert_image(page.rect,filename=str(image));document.save(pdf)
        doc=Document();doc.add_picture(str(image));doc.save(root/'synthetic.docx')
        deck=Presentation();slide=deck.slides.add_slide(deck.slide_layouts[6]);slide.shapes.add_picture(str(image),Inches(1),Inches(1));deck.save(root/'synthetic.pptx')
        for path in (image,pdf,root/'synthetic.docx',root/'synthetic.pptx'):
            text=extract_local_ocr(path,backend='rapidocr',timeout=60)
            if INCOMPLETE in text or not all(word in text.upper() for word in ('COURSE','QUESTIONS')):
                raise RuntimeError('Local OCR smoke failed for '+path.suffix)
            print('PASS local OCR '+path.suffix+' (synthetic image text recognized)')


if __name__=='__main__':main()
